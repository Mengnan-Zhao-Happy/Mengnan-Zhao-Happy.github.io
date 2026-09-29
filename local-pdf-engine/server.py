from __future__ import annotations

import base64
import io
import json
import os
import re
import tempfile
from pathlib import Path
from typing import Iterable

import fitz
import uvicorn
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from starlette.background import BackgroundTask
from formula_extractor import MATH_TOKEN, block_text, font_is_math, is_label_block, is_math_display
from latex2mathml.converter import convert as latex_to_mathml
from omml_engine import latex_to_oMath

os.environ.setdefault("NO_ALBUMENTATIONS_UPDATE", "1")

app = FastAPI(title="Local PDF to Editable Word Engine")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "https://advlearnlab.github.io",
        "https://mengnan-zhao-happy.github.io",
        "http://127.0.0.1:4321",
        "http://localhost:4321",
    ],
    allow_methods=["GET", "POST"],
    allow_headers=["*"],
    expose_headers=["X-PDF-Formulas"],
)


@app.middleware("http")
async def allow_private_network(request, call_next):
    response = await call_next(request)
    if request.headers.get("access-control-request-private-network") == "true":
        response.headers["Access-Control-Allow-Private-Network"] = "true"
    return response

_formula_model = None
_control_chars = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]")
_math_fonts = re.compile(r"(cmmi|cmsy|cmex|math|symbol|mt extra|euclid)", re.I)
_math_symbols = re.compile(r"[=<>≤≥≠≈∑∏∫√±×÷∞∂∇∈⊂^_{}]|[α-ωΑ-Ω]")


def clean_text(value: str) -> str:
    return _control_chars.sub("", value or "").replace("\u00ad", "").strip()


def parse_pages(spec: str, total: int) -> list[int]:
    if not spec.strip():
        return list(range(total))
    selected: set[int] = set()
    for part in re.split(r"[,，\s]+", spec.strip()):
        match = re.fullmatch(r"(\d+)(?:-(\d+))?", part)
        if not match:
            continue
        start, end = int(match.group(1)), int(match.group(2) or match.group(1))
        start, end = sorted((start, end))
        selected.update(range(max(1, start) - 1, min(total, end)))
    return sorted(selected)


def block_order(blocks: list[dict], page_width: float) -> list[dict]:
    usable = [block for block in blocks if block.get("type") in (0, 1)]
    midpoint = page_width / 2
    narrow = [block for block in usable if block["bbox"][2] - block["bbox"][0] < page_width * 0.72]
    left_count = sum(1 for block in narrow if block["bbox"][0] < midpoint)
    right_count = sum(1 for block in narrow if block["bbox"][0] >= midpoint)
    if left_count < 3 or right_count < 3:
        return sorted(usable, key=lambda block: (block["bbox"][1], block["bbox"][0]))
    top = min((block["bbox"][1] for block in narrow), default=0)
    preamble = [block for block in usable if block["bbox"][1] < top]
    remaining = [block for block in usable if block not in preamble]
    left = [block for block in remaining if block["bbox"][0] < midpoint]
    right = [block for block in remaining if block["bbox"][0] >= midpoint]
    key = lambda block: (block["bbox"][1], block["bbox"][0])
    return sorted(preamble, key=key) + sorted(left, key=key) + sorted(right, key=key)


def line_text(line: dict) -> str:
    return clean_text("".join(span.get("text", "") for span in line.get("spans", [])))


def is_display_formula(block: dict, page_width: float) -> bool:
    if is_math_display(block, page_width):
        return True
    text = block_text(block)
    if not text or is_label_block(block, page_width) or len(text) > 500:
        return False
    spans = [span for line in block.get("lines", []) for span in line.get("spans", []) if clean_text(span.get("text", ""))]
    if not spans or not MATH_TOKEN.search(text):
        return False
    math_chars = sum(len(clean_text(span.get("text", ""))) for span in spans if font_is_math(span.get("font", "")))
    total_chars = sum(len(clean_text(span.get("text", ""))) for span in spans)
    prose_words = len(re.findall(r"[A-Za-z]{5,}", text))
    return math_chars / max(1, total_chars) >= 0.42 and prose_words <= 5


def merge_display_formula_blocks(blocks: list[dict], page_width: float) -> list[dict]:
    merged: list[dict] = []
    index = 0
    while index < len(blocks):
        block = blocks[index]
        if block.get("type") != 0 or not is_display_formula(block, page_width):
            merged.append(block); index += 1; continue
        group = [block]
        cursor = index + 1
        bottom = block["bbox"][3]
        while cursor < len(blocks):
            candidate = blocks[cursor]
            if candidate.get("type") != 0 or candidate["bbox"][1] - bottom > 18:
                break
            has_math_font = any(font_is_math(span.get("font", "")) for line in candidate.get("lines", []) for span in line.get("spans", []))
            if not (is_display_formula(candidate, page_width) or is_label_block(candidate, page_width) or has_math_font):
                break
            group.append(candidate); bottom = max(bottom, candidate["bbox"][3]); cursor += 1
        if len(group) == 1:
            merged.append(block)
        else:
            bbox = (min(item["bbox"][0] for item in group), min(item["bbox"][1] for item in group), max(item["bbox"][2] for item in group), max(item["bbox"][3] for item in group))
            merged.append({"type": 0, "bbox": bbox, "lines": [line for item in group for line in item.get("lines", [])], "formula_group": True})
        index = cursor
    return merged


def is_formula_line(line: dict) -> bool:
    text = line_text(line)
    if len(text) < 2:
        return False
    spans = line.get("spans", [])
    math_font = any(_math_fonts.search(span.get("font", "")) for span in spans)
    symbol_count = len(_math_symbols.findall(text))
    word_count = len(re.findall(r"[A-Za-z]{5,}", text))
    return (math_font and symbol_count >= 1) or (symbol_count >= 2 and word_count <= 4)


def join_lines(lines: Iterable[str]) -> str:
    result = ""
    for value in lines:
        value = clean_text(value)
        if not value:
            continue
        if result.endswith("-") and value[:1].islower():
            result = result[:-1] + value
        else:
            result += (" " if result else "") + value
    return re.sub(r"\s+", " ", result).strip()


def formula_model():
    global _formula_model
    if _formula_model is None:
        from pix2tex.cli import LatexOCR

        _formula_model = LatexOCR()
    return _formula_model


def formula_weights_ready() -> bool:
    try:
        import pix2tex
        package_root = Path(pix2tex.__file__).resolve().parent
        weights = package_root / "model" / "checkpoints" / "weights.pth"
        resizer = package_root / "model" / "checkpoints" / "image_resizer.pth"
        return weights.exists() and weights.stat().st_size == 102_113_875 and resizer.exists() and resizer.stat().st_size == 19_441_973
    except (ImportError, OSError):
        return False


def embedded_formula(text: str) -> str:
    replacements = {
        "α": r"\alpha", "β": r"\beta", "γ": r"\gamma", "δ": r"\delta",
        "θ": r"\theta", "λ": r"\lambda", "μ": r"\mu", "π": r"\pi",
        "σ": r"\sigma", "φ": r"\phi", "ω": r"\omega", "×": r"\times",
        "±": r"\pm", "≤": r"\leq", "≥": r"\geq", "≠": r"\neq",
        "≈": r"\approx", "∞": r"\infty", "∑": r"\sum", "∏": r"\prod",
        "∫": r"\int", "∂": r"\partial", "∇": r"\nabla", "∈": r"\in",
    }
    value = clean_text(text)
    for symbol, command in replacements.items():
        value = value.replace(symbol, f" {command} ")
    return re.sub(r"\s+", " ", value).strip()


def normalize_latex(value: str) -> str:
    value = clean_text(value).strip()
    value = re.sub(r"^```(?:latex)?|```$", "", value, flags=re.I).strip()
    value = re.sub(r"^(?:\$\$|\\\[|\\\()|(?:\$\$|\\\]|\\\))$", "", value).strip()
    value = re.sub(r"\\tag\s*\{[^{}]*\}\s*$", "", value).strip()
    value = re.sub(r"\s*\(\s*\d+(?:\.\d+)*\s*\)\s*$", "", value).strip()
    return re.sub(r"\s+", " ", value)


def valid_latex(candidate: str, source: str) -> bool:
    if not candidate or len(candidate) > 600 or re.search(r"(.)\1{10,}", candidate):
        return False
    if candidate.count("{") != candidate.count("}"):
        return False
    if any(token in candidate for token in (r"\includegraphics", r"\documentclass", r"\begin{document}")):
        return False
    if "=" in source and "=" not in candidate:
        return False
    source_ops = set(re.findall(r"[=<>≤≥≠≈∑∏∫√±×÷∞]", source))
    candidate_ops = set(re.findall(r"[=<>≤≥≠≈∑∏∫√±×÷∞]", candidate))
    if source_ops and not candidate_ops and not re.search(r"\\(?:sum|prod|int|sqrt|leq|geq|neq|approx|pm|times)", candidate):
        return False
    try:
        latex_to_mathml(candidate)
    except Exception:
        return False
    return True


def recognize_formula(page: fitz.Page, bbox: tuple[float, float, float, float], source: str) -> str:
    if not formula_weights_ready():
        return ""
    rect = fitz.Rect(bbox)
    rect.x0 = max(page.rect.x0, rect.x0 - 8)
    rect.y0 = max(page.rect.y0, rect.y0 - 6)
    rect.x1 = min(page.rect.x1, rect.x1 + 8)
    rect.y1 = min(page.rect.y1, rect.y1 + 6)
    pixmap = page.get_pixmap(matrix=fitz.Matrix(4.2, 4.2), clip=rect, alpha=False)
    from PIL import Image, ImageOps

    image = Image.open(io.BytesIO(pixmap.tobytes("png"))).convert("L")
    image = ImageOps.expand(ImageOps.autocontrast(image), border=24, fill="white").convert("RGB")
    try:
        latex = normalize_latex(formula_model()(image))
    except (ImportError, ModuleNotFoundError, RuntimeError):
        return ""
    return latex if valid_latex(latex, source) else ""


def omml_run(text: str) -> str:
    escaped = (
        text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    )
    return f'<m:r><m:t xml:space="preserve">{escaped}</m:t></m:r>'


def latex_to_omml(latex: str) -> str:
    command_map = {
        "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε",
        "theta": "θ", "lambda": "λ", "mu": "μ", "pi": "π", "rho": "ρ", "sigma": "σ",
        "tau": "τ", "phi": "φ", "omega": "ω", "times": "×", "cdot": "·", "pm": "±",
        "leq": "≤", "geq": "≥", "neq": "≠", "approx": "≈", "infty": "∞", "sum": "∑",
        "prod": "∏", "int": "∫", "partial": "∂", "nabla": "∇", "in": "∈", "subset": "⊂",
    }
    source = re.sub(r"\\(?:left|right)", "", latex)
    cursor = 0

    def sequence(stop: str = "") -> str:
        nonlocal cursor
        output = ""
        while cursor < len(source) and (not stop or source[cursor] != stop):
            output += atom()
        if stop and cursor < len(source) and source[cursor] == stop:
            cursor += 1
        return output

    def argument() -> str:
        nonlocal cursor
        if cursor < len(source) and source[cursor] == "{":
            cursor += 1
            return sequence("}")
        return atom(False)

    def atom(scripts: bool = True) -> str:
        nonlocal cursor
        if cursor >= len(source):
            return ""
        if source[cursor] == "{":
            cursor += 1
            base = sequence("}")
        elif source[cursor] == "\\":
            cursor += 1
            match = re.match(r"[A-Za-z]+", source[cursor:])
            command = match.group(0) if match else source[cursor:cursor + 1]
            cursor += len(command)
            if command == "frac":
                base = f"<m:f><m:num>{argument()}</m:num><m:den>{argument()}</m:den></m:f>"
            elif command == "sqrt":
                base = f'<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:e>{argument()}</m:e></m:rad>'
            elif command in ("text", "mathrm", "mathbf", "mathit"):
                base = argument()
            else:
                base = omml_run(command_map.get(command, "\\" + command))
        else:
            base = omml_run(source[cursor])
            cursor += 1
        if not scripts:
            return base
        sub = sup = ""
        while cursor < len(source) and source[cursor] in "_^":
            kind = source[cursor]
            cursor += 1
            value = argument()
            if kind == "_": sub = value
            else: sup = value
        if sub and sup:
            return f"<m:sSubSup><m:e>{base}</m:e><m:sub>{sub}</m:sub><m:sup>{sup}</m:sup></m:sSubSup>"
        if sub:
            return f"<m:sSub><m:e>{base}</m:e><m:sub>{sub}</m:sub></m:sSub>"
        if sup:
            return f"<m:sSup><m:e>{base}</m:e><m:sup>{sup}</m:sup></m:sSup>"
        return base

    return f'<m:oMath {nsdecls("m")}>{sequence()}</m:oMath>'


def equation_element(latex: str, display: bool):
    element = latex_to_oMath(latex)
    if not display:
        return element
    from lxml import etree
    wrapper = etree.Element("{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara")
    wrapper.append(element)
    return wrapper


def add_equation(document: Document, latex: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = 1
    paragraph._p.append(equation_element(latex, display=True))


def add_formula_image(document: Document, page: fitz.Page, bbox, page_width: float) -> None:
    rect = fitz.Rect(bbox)
    rect.x0 = max(page.rect.x0, rect.x0 - 6)
    rect.y0 = max(page.rect.y0, rect.y0 - 4)
    rect.x1 = min(page.rect.x1, rect.x1 + 6)
    rect.y1 = min(page.rect.y1, rect.y1 + 4)
    pixmap = page.get_pixmap(matrix=fitz.Matrix(3, 3), clip=rect, alpha=False)
    paragraph = document.add_paragraph(); paragraph.alignment = 1
    width = min(6.0, max(1.0, rect.width / page_width * 6.2))
    paragraph.add_run().add_picture(io.BytesIO(pixmap.tobytes("png")), width=Inches(width))


def add_text_block(document: Document, block: dict, page: fitz.Page, page_number: int, formulas: list[dict]) -> None:
    lines = block.get("lines", [])
    source = block_text(block)
    if is_display_formula(block, page.rect.width):
        latex = recognize_formula(page, block["bbox"], source)
        method = "OCR"
        if not latex:
            fallback = embedded_formula(source)
            latex = fallback if valid_latex(fallback, source) else ""
            method = "PDF text"
        if latex:
            try:
                add_equation(document, latex)
                formulas.append({"pageNumber": page_number, "latex": latex, "method": method})
                return
            except Exception:
                pass
        add_formula_image(document, page, block["bbox"], page.rect.width)
        formulas.append({"pageNumber": page_number, "latex": "", "method": "image fallback"})
        return

    text = join_lines(line_text(line) for line in lines)
    if not text:
        return
    max_size = max((span.get("size", 10) for line in lines for span in line.get("spans", [])), default=10)
    if len(text) < 120 and max_size >= 14:
        document.add_heading(text, level=1 if max_size >= 18 else 2)
        return
    paragraph = document.add_paragraph()
    for line_index, line in enumerate(lines):
        math_buffer = ""
        for span in line.get("spans", []):
            span_text = clean_text(span.get("text", ""))
            is_math = font_is_math(span.get("font", "")) and bool(MATH_TOKEN.search(span_text))
            if is_math:
                math_buffer += span_text
                continue
            if math_buffer:
                latex = embedded_formula(math_buffer)
                try:
                    paragraph._p.append(equation_element(latex, display=False))
                    formulas.append({"pageNumber": page_number, "latex": latex, "method": "PDF inline"})
                except Exception:
                    paragraph.add_run(math_buffer)
                math_buffer = ""
            paragraph.add_run(span_text)
        if math_buffer:
            latex = embedded_formula(math_buffer)
            try:
                paragraph._p.append(equation_element(latex, display=False))
                formulas.append({"pageNumber": page_number, "latex": latex, "method": "PDF inline"})
            except Exception:
                paragraph.add_run(math_buffer)
        if line_index < len(lines) - 1:
            paragraph.add_run(" ")


def add_image_block(document: Document, block: dict, page_width: float) -> None:
    image = block.get("image")
    if not image:
        return
    bbox = block["bbox"]
    width_inches = min(6.2, max(1.0, (bbox[2] - bbox[0]) / page_width * 6.2))
    try:
        paragraph = document.add_paragraph()
        paragraph.alignment = 1
        paragraph.add_run().add_picture(io.BytesIO(image), width=Inches(width_inches))
    except Exception:
        pass


def convert_pdf(source: Path, destination: Path, pages_spec: str) -> list[dict]:
    pdf = fitz.open(source)
    selected = parse_pages(pages_spec, len(pdf))
    if not selected:
        raise ValueError("No valid pages selected")
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.72)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    formulas: list[dict] = []
    for page_index, page_number in enumerate(selected):
        page = pdf[page_number]
        payload = page.get_text("dict", flags=fitz.TEXTFLAGS_DICT)
        ordered_blocks = block_order(payload.get("blocks", []), page.rect.width)
        for block in merge_display_formula_blocks(ordered_blocks, page.rect.width):
            if block.get("type") == 0:
                add_text_block(document, block, page, page_number + 1, formulas)
            elif block.get("type") == 1:
                add_image_block(document, block, page.rect.width)
        if page_index < len(selected) - 1:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.save(destination)
    pdf.close()
    return formulas


@app.get("/status")
def status():
    formula_available = formula_weights_ready()
    engine = "PyMuPDF + OOXML" + (" + pix2tex" if formula_available else "")
    return {"ready": True, "formulaModelInstalled": formula_available, "formulaModelLoaded": _formula_model is not None, "engine": engine}


@app.post("/convert")
async def convert(file: UploadFile = File(...), pages: str = Form("")):
    if not file.filename or not file.filename.lower().endswith(".pdf"):
        raise HTTPException(400, "Please upload a PDF file")
    directory = Path(tempfile.mkdtemp(prefix="pdf_word_"))
    source = directory / "source.pdf"
    destination = directory / "converted-editable.docx"
    source.write_bytes(await file.read())
    try:
        formulas = convert_pdf(source, destination, pages)
    except Exception as error:
        import shutil
        shutil.rmtree(directory, ignore_errors=True)
        raise HTTPException(500, str(error)) from error
    output_name = Path(file.filename).stem + "-editable.docx"
    formula_json = json.dumps(formulas[:40], ensure_ascii=False).encode("utf-8")
    headers = {"X-PDF-Formulas": base64.b64encode(formula_json).decode("ascii")}
    return FileResponse(destination, filename=output_name, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers=headers, background=BackgroundTask(lambda: __import__("shutil").rmtree(directory, ignore_errors=True)))


if __name__ == "__main__":
    if formula_weights_ready():
        formula_model()
    uvicorn.run(app, host="127.0.0.1", port=8765)
